import pandas as pd
import os
import sys
from docx import Document

# Fix encoding for Windows
if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8')

# 1. Read EHR from Excel
print("=" * 60)
print("EXTRACTING EHR DATA FROM EXCEL")
print("=" * 60)
ehr_file = r'C:\Users\Kashish\Downloads\ehr.xlsx'
ehr_df = pd.read_excel(ehr_file)
print("EHR Data Preview:")
print(ehr_df.to_string())
print(f"\nShape: {ehr_df.shape}")
print(f"Columns: {list(ehr_df.columns)}\n")

# Save to bronze
ehr_df.to_csv('bronze/ehr.csv', index=False)
print("[OK] Saved to bronze/ehr.csv")

# 2. Read Vitals from Word (JSON format)
print("\n" + "=" * 60)
print("EXTRACTING VITALS DATA FROM WORD")
print("=" * 60)
vitals_file = r'C:\Users\Kashish\Downloads\vitals.docx'
try:
    doc = Document(vitals_file)
    vitals_list = []
    for para in doc.paragraphs:
        if para.text.strip() and para.text.strip().startswith('{'):
            try:
                import json
                vital = json.loads(para.text)
                vitals_list.append(vital)
            except:
                pass
    
    if vitals_list:
        vitals_df = pd.DataFrame(vitals_list)
        # Rename columns to match expected format
        if 'patientId' in vitals_df.columns:
            vitals_df.rename(columns={'patientId': 'patient_id'}, inplace=True)
        print(vitals_df.head().to_string())
        print(f"... ({len(vitals_df)} rows total)")
        vitals_df.to_csv('bronze/vitals.csv', index=False)
        print("[OK] Saved to bronze/vitals.csv\n")
    else:
        print("[ERROR] No vitals data found")
except Exception as e:
    print(f"[ERROR] Error reading vitals.docx: {e}")

# 3. Read Labs from Word (JSON format)
print("\n" + "=" * 60)
print("EXTRACTING LABS DATA FROM WORD")
print("=" * 60)
labs_file = r'C:\Users\Kashish\Downloads\labs- Ajay kumar.docx'
try:
    doc = Document(labs_file)
    labs_text = ""
    for para in doc.paragraphs:
        labs_text += para.text + "\n"
    
    # Try to parse JSON
    import json
    labs_text = labs_text.strip()
    if labs_text.startswith('['):
        labs_list = json.loads(labs_text)
    else:
        # Try to extract JSON array from text
        start = labs_text.find('[')
        end = labs_text.rfind(']') + 1
        if start >= 0 and end > start:
            labs_list = json.loads(labs_text[start:end])
        else:
            labs_list = []
    
    if labs_list:
        labs_df = pd.DataFrame(labs_list)
        print(labs_df.head().to_string())
        print(f"... ({len(labs_df)} rows total)")
        labs_df.to_csv('bronze/labs.csv', index=False)
        print("[OK] Saved to bronze/labs.csv\n")
    else:
        print("[ERROR] No labs data found")
except Exception as e:
    print(f"[ERROR] Error reading labs docx: {e}")

print("=" * 60)
print("[SUCCESS] DATA EXTRACTION COMPLETE")
print("=" * 60)
