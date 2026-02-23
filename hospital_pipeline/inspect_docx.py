"""
Inspect the structure of .docx files
"""
from docx import Document

print("=" * 80)
print("Inspecting vitals.docx")
print("=" * 80)
try:
    doc = Document('vitals.docx')
    print(f"Number of paragraphs: {len(doc.paragraphs)}")
    print(f"Number of tables: {len(doc.tables)}")
    
    if doc.paragraphs:
        print("\nFirst 10 paragraphs:")
        for i, para in enumerate(doc.paragraphs[:10]):
            if para.text.strip():
                print(f"  {i+1}. {para.text[:100]}")
    
    if doc.tables:
        print(f"\nTable structure:")
        for idx, table in enumerate(doc.tables):
            print(f"\n  Table {idx+1}:")
            print(f"    Rows: {len(table.rows)}")
            print(f"    Columns: {len(table.rows[0].cells) if table.rows else 0}")
            if table.rows:
                print(f"    First row: {[cell.text.strip()[:30] for cell in table.rows[0].cells]}")
except Exception as e:
    print(f"Error reading vitals.docx: {e}")

print("\n" + "=" * 80)
print("Inspecting labs.docx")
print("=" * 80)
try:
    doc = Document('labs.docx')
    print(f"Number of paragraphs: {len(doc.paragraphs)}")
    print(f"Number of tables: {len(doc.tables)}")
    
    if doc.paragraphs:
        print("\nFirst 10 paragraphs:")
        for i, para in enumerate(doc.paragraphs[:10]):
            if para.text.strip():
                print(f"  {i+1}. {para.text[:100]}")
    
    if doc.tables:
        print(f"\nTable structure:")
        for idx, table in enumerate(doc.tables):
            print(f"\n  Table {idx+1}:")
            print(f"    Rows: {len(table.rows)}")
            print(f"    Columns: {len(table.rows[0].cells) if table.rows else 0}")
            if table.rows:
                print(f"    First row: {[cell.text.strip()[:30] for cell in table.rows[0].cells]}")
except Exception as e:
    print(f"Error reading labs.docx: {e}")
