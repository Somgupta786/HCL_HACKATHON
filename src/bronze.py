"""Bronze layer - Raw data ingestion."""
import pandas as pd
import json
from docx import Document


def ingest_ehr():
    """Read EHR Excel and save to bronze layer."""
    df = pd.read_excel('data/ehr.xlsx')
    df.to_csv('bronze/ehr.csv', index=False)


def ingest_vitals():
    """Read vitals DOCX and save to bronze layer."""
    doc = Document('data/vitals.docx')
    records = []
    for para in doc.paragraphs:
        if para.text.strip():
            records.append(json.loads(para.text))
    df = pd.DataFrame(records)
    df.to_csv('bronze/vitals.csv', index=False)


def ingest_labs():
    """Read labs DOCX and save to bronze layer."""
    doc = Document('data/labs.docx')
    full_text = '\n'.join([para.text for para in doc.paragraphs])
    data = json.loads(full_text)
    df = pd.DataFrame(data)
    df.to_csv('bronze/labs.csv', index=False)


def run_bronze():
    """Execute all bronze layer ingestion."""
    ingest_ehr()
    ingest_vitals()
    ingest_labs()
