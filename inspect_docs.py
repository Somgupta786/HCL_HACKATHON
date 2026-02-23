from docx import Document

# Check vitals.docx
print("=== VITALS.DOCX ===")
doc = Document('data/vitals.docx')
for i, para in enumerate(doc.paragraphs[:10]):
    if para.text.strip():
        print(f"Line {i}: {para.text}")

print("\n=== LABS.DOCX ===")
doc = Document('data/labs.docx')
for i, para in enumerate(doc.paragraphs[:10]):
    if para.text.strip():
        print(f"Line {i}: {para.text}")
