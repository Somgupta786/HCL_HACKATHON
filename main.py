"""
Hospital Health Monitoring Pipeline
Main orchestration script for Bronze → Silver → Gold data processing
"""

import pandas as pd
import numpy as np
import json
import os
from pathlib import Path
from datetime import datetime
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from docx import Document
import openpyxl


def create_directories():
    """Create required directories if they don't exist"""
    directories = ['bronze', 'silver', 'gold', 'visualizations']
    for directory in directories:
        Path(directory).mkdir(exist_ok=True)
    print("✓ Directories created/verified")


# ====================================================================================
# STEP 1: BRONZE LAYER - RAW STORAGE
# ====================================================================================

def load_to_bronze():
    """
    Load all files from data/ to bronze/ as CSV
    Handles multiple formats: CSV, Excel, JSON, JSONL, Word docs
    No cleaning or transformations - raw storage only
    """
    print("\n" + "="*70)
    print("STEP 1: BRONZE LAYER - RAW STORAGE")
    print("="*70)
    
    data_dir = Path('data')
    
    # 1. Load EHR data (try multiple formats)
    print("\n[1/3] Loading EHR data...")
    ehr_loaded = False
    
    # Try Excel first
    if (data_dir / 'ehr.xlsx').exists():
        df_ehr = pd.read_excel(data_dir / 'ehr.xlsx')
        ehr_loaded = True
    # Try CSV
    elif (data_dir / 'ehr.csv').exists():
        df_ehr = pd.read_csv(data_dir / 'ehr.csv')
        ehr_loaded = True
    # Try other Excel formats
    elif (data_dir / 'ehr.xls').exists():
        df_ehr = pd.read_excel(data_dir / 'ehr.xls')
        ehr_loaded = True
    
    if ehr_loaded:
        df_ehr.to_csv('bronze/ehr.csv', index=False)
        print(f"  ✓ Saved bronze/ehr.csv ({len(df_ehr)} records)")
    else:
        print(f"  ⚠ No EHR file found (tried: ehr.xlsx, ehr.csv, ehr.xls)")
    
    # 2. Load Vitals data (try multiple formats)
    print("\n[2/3] Loading Vitals data...")
    vitals_loaded = False
    
    # Try JSONL
    if (data_dir / 'vitals.jsonl').exists():
        vitals_records = []
        with open(data_dir / 'vitals.jsonl', 'r') as f:
            for line in f:
                vitals_records.append(json.loads(line))
        df_vitals = pd.DataFrame(vitals_records)
        vitals_loaded = True
    # Try JSON
    elif (data_dir / 'vitals.json').exists():
        with open(data_dir / 'vitals.json', 'r') as f:
            vitals_data = json.load(f)
        df_vitals = pd.DataFrame(vitals_data)
        vitals_loaded = True
    # Try CSV
    elif (data_dir / 'vitals.csv').exists():
        df_vitals = pd.read_csv(data_dir / 'vitals.csv')
        vitals_loaded = True
    # Try Word doc (extract JSON from paragraphs)
    elif (data_dir / 'vitals.docx').exists():
        try:
            doc = Document(data_dir / 'vitals.docx')
            vitals_records = []
            
            # Extract JSON objects from paragraphs
            for para in doc.paragraphs:
                text = para.text.strip()
                if text and (text.startswith('{') or text.startswith('[')):
                    try:
                        # Try to parse as JSON
                        data = json.loads(text)
                        if isinstance(data, dict):
                            vitals_records.append(data)
                        elif isinstance(data, list):
                            vitals_records.extend(data)
                    except json.JSONDecodeError:
                        continue
            
            if vitals_records:
                df_vitals = pd.DataFrame(vitals_records)
                vitals_loaded = True
                print(f"  → Extracted {len(vitals_records)} records from vitals.docx")
            else:
                print(f"  → Warning: No JSON data found in vitals.docx")
        except Exception as e:
            print(f"  → Error reading vitals.docx: {str(e)}")
    
    if vitals_loaded:
        df_vitals.to_csv('bronze/vitals.csv', index=False)
        print(f"  ✓ Saved bronze/vitals.csv ({len(df_vitals)} records)")
    else:
        print(f"  ⚠ No Vitals file found (tried: vitals.jsonl, vitals.json, vitals.csv, vitals.docx)")
    
    # 3. Load Labs data (try multiple formats)
    print("\n[3/3] Loading Labs data...")
    labs_loaded = False
    
    # Try JSON
    if (data_dir / 'labs.json').exists():
        with open(data_dir / 'labs.json', 'r') as f:
            labs_data = json.load(f)
        df_labs = pd.DataFrame(labs_data)
        labs_loaded = True
    # Try CSV
    elif (data_dir / 'labs.csv').exists():
        df_labs = pd.read_csv(data_dir / 'labs.csv')
        labs_loaded = True
    # Try Word doc (extract JSON from paragraphs) - try both filenames
    else:
        # Check for any labs*.docx file
        labs_files = list(data_dir.glob('labs*.docx'))
        if labs_files:
            labs_file = labs_files[0]  # Use first match
            try:
                doc = Document(labs_file)
                labs_records = []
                
                # Collect all paragraphs to handle JSON array split across lines
                full_text = ""
                for para in doc.paragraphs:
                    text = para.text.strip()
                    if text:
                        full_text += text + "\n"
                
                # Try to parse as complete JSON array
                try:
                    data = json.loads(full_text)
                    if isinstance(data, list):
                        labs_records = data
                    elif isinstance(data, dict):
                        labs_records = [data]
                except json.JSONDecodeError:
                    # If full text parsing fails, try line by line
                    for para in doc.paragraphs:
                        text = para.text.strip()
                        if text and (text.startswith('{') or text.startswith('[')):
                            try:
                                data = json.loads(text)
                                if isinstance(data, dict):
                                    labs_records.append(data)
                                elif isinstance(data, list):
                                    labs_records.extend(data)
                            except json.JSONDecodeError:
                                continue
                
                if labs_records:
                    df_labs = pd.DataFrame(labs_records)
                    labs_loaded = True
                    print(f"  → Extracted {len(labs_records)} records from {labs_file.name}")
                else:
                    print(f"  → Warning: No JSON data found in {labs_file.name}")
            except Exception as e:
                print(f"  → Error reading {labs_file.name}: {str(e)}")
    
    if labs_loaded:
        df_labs.to_csv('bronze/labs.csv', index=False)
        print(f"  ✓ Saved bronze/labs.csv ({len(df_labs)} records)")
    else:
        print(f"  ⚠ No Labs file found (tried: labs.json, labs.csv, labs*.docx)")
    
    print("\n✅ Bronze layer complete - all raw data stored")


# ====================================================================================
# STEP 2: SILVER LAYER - CLEANING & STANDARDIZATION
# ====================================================================================

def clean_vitals():
    """
    Clean vitals data:
    - Convert UNIX timestamp to datetime
    - Ensure numeric columns
    - Standardize column names
    """
    print("\n[1/2] Cleaning Vitals data...")
    
    df = pd.read_csv('bronze/vitals.csv')
    
    # Standardize column names
    df = df.rename(columns={
        'patientId': 'patient_id'
    })
    
    # Convert UNIX timestamp to datetime
    df['timestamp'] = pd.to_datetime(df['timestamp'], unit='s')
    
    # Ensure numeric columns
    numeric_cols = ['hr', 'ox', 'sys', 'dia']
    for col in numeric_cols:
        df[col] = pd.to_numeric(df[col], errors='coerce')
    
    # Save cleaned data
    df.to_csv('silver/clean_vitals.csv', index=False)
    print(f"  ✓ Saved silver/clean_vitals.csv ({len(df)} records)")
    
    return df


def clean_labs():
    """
    Clean labs data:
    - Convert timestamp to datetime
    - Ensure numeric value column
    - Standardize column names
    """
    print("\n[2/2] Cleaning Labs data...")
    
    df = pd.read_csv('bronze/labs.csv')
    
    # Standardize column names
    df = df.rename(columns={
        'patientId': 'patient_id',
        'test': 'lab_test',
        'value': 'lab_value'
    })
    
    # Convert timestamp to datetime
    df['timestamp'] = pd.to_datetime(df['timestamp'])
    
    # Ensure numeric lab_value
    df['lab_value'] = pd.to_numeric(df['lab_value'], errors='coerce')
    
    # Save cleaned data
    df.to_csv('silver/clean_labs.csv', index=False)
    print(f"  ✓ Saved silver/clean_labs.csv ({len(df)} records)")
    
    return df


def clean_silver_layer():
    """Execute all silver layer cleaning"""
    print("\n" + "="*70)
    print("STEP 2: SILVER LAYER - CLEANING & STANDARDIZATION")
    print("="*70)
    
    df_vitals = clean_vitals()
    df_labs = clean_labs()
    
    print("\n✅ Silver layer complete - all data cleaned and standardized")
    
    return df_vitals, df_labs


# ====================================================================================
# STEP 3: PATIENT MASTER TABLE
# ====================================================================================

def create_patient_master(df_vitals, df_labs):
    """
    Create unified patient master table
    Join: EHR + Latest Vitals + Latest Labs
    """
    print("\n" + "="*70)
    print("STEP 3: PATIENT MASTER TABLE")
    print("="*70)
    
    # Load EHR
    df_ehr = pd.read_csv('bronze/ehr.csv')
    print(f"\n[1/3] Loaded EHR: {len(df_ehr)} patients")
    
    # Get latest vitals per patient (max timestamp)
    latest_vitals = df_vitals.loc[df_vitals.groupby('patient_id')['timestamp'].idxmax()]
    latest_vitals = latest_vitals.add_prefix('latest_vitals_')
    latest_vitals = latest_vitals.rename(columns={'latest_vitals_patient_id': 'patient_id'})
    print(f"[2/3] Latest vitals extracted: {len(latest_vitals)} patients")
    
    # Get latest lab values per patient per test
    latest_labs = df_labs.loc[df_labs.groupby(['patient_id', 'lab_test'])['timestamp'].idxmax()]
    latest_labs_pivot = latest_labs.pivot(
        index='patient_id',
        columns='lab_test',
        values='lab_value'
    ).reset_index()
    latest_labs_pivot.columns = ['patient_id'] + [f'latest_{col}' for col in latest_labs_pivot.columns[1:]]
    print(f"[3/3] Latest labs extracted: {len(latest_labs_pivot)} patients")
    
    # Join all together
    patient_master = df_ehr.merge(latest_vitals, on='patient_id', how='left')
    patient_master = patient_master.merge(latest_labs_pivot, on='patient_id', how='left')
    
    # Save patient master
    patient_master.to_csv('silver/patient_master.csv', index=False)
    print(f"\n✓ Saved silver/patient_master.csv ({len(patient_master)} patients)")
    print(f"  Columns: {', '.join(patient_master.columns)}")
    
    print("\n✅ Patient master table created")
    
    return patient_master


# ====================================================================================
# STEP 4: ANOMALY DETECTION (GOLD LAYER)
# ====================================================================================

def detect_anomalies(df_vitals):
    """
    Rule-based anomaly detection:
    - HR > 120 → High Heart Rate
    - OX < 92 → Low Oxygen
    - SYS > 160 OR DIA > 100 → High Blood Pressure
    """
    print("\n" + "="*70)
    print("STEP 4: ANOMALY DETECTION (GOLD LAYER)")
    print("="*70)
    
    anomalies = []
    
    for _, row in df_vitals.iterrows():
        patient_id = row['patient_id']
        timestamp = row['timestamp']
        
        # Rule 1: High Heart Rate
        if pd.notna(row['hr']) and row['hr'] > 120:
            anomalies.append({
                'patient_id': patient_id,
                'timestamp': timestamp,
                'anomaly': 'High Heart Rate',
                'value': row['hr']
            })
        
        # Rule 2: Low Oxygen
        if pd.notna(row['ox']) and row['ox'] < 92:
            anomalies.append({
                'patient_id': patient_id,
                'timestamp': timestamp,
                'anomaly': 'Low Oxygen',
                'value': row['ox']
            })
        
        # Rule 3: High Blood Pressure
        if (pd.notna(row['sys']) and row['sys'] > 160) or (pd.notna(row['dia']) and row['dia'] > 100):
            anomalies.append({
                'patient_id': patient_id,
                'timestamp': timestamp,
                'anomaly': 'High Blood Pressure',
                'value': f"{row['sys']}/{row['dia']}"
            })
    
    df_anomalies = pd.DataFrame(anomalies)
    df_anomalies.to_csv('gold/anomalies.csv', index=False)
    
    print(f"\n✓ Saved gold/anomalies.csv ({len(anomalies)} anomalies detected)")
    
    if len(anomalies) > 0:
        print("\nAnomaly breakdown:")
        for anomaly_type, count in df_anomalies['anomaly'].value_counts().items():
            print(f"  - {anomaly_type}: {count}")
    
    print("\n✅ Anomaly detection complete")
    
    return df_anomalies


# ====================================================================================
# STEP 5: VISUALIZATIONS
# ====================================================================================

def generate_visualizations(df_vitals, df_anomalies):
    """Generate all required visualizations and save as PNG"""
    print("\n" + "="*70)
    print("STEP 5: VISUALIZATIONS")
    print("="*70)
    
    # Set matplotlib style
    plt.style.use('seaborn-v0_8-darkgrid')
    
    # ----- Chart 1: Heart Rate Trend -----
    print("\n[1/3] Generating Heart Rate Trend chart...")
    
    fig, ax = plt.subplots(figsize=(14, 6))
    
    # Calculate statistics over time for cleaner visualization
    stats_by_time = df_vitals.groupby('timestamp')['hr'].agg(['mean', 'min', 'max', 'std']).reset_index()
    stats_by_time = stats_by_time.sort_values('timestamp')
    
    # Plot min-max range as filled area
    ax.fill_between(stats_by_time['timestamp'], stats_by_time['min'], stats_by_time['max'],
                     alpha=0.2, color='#3498DB', label='Min-Max Range')
    
    # Plot standard deviation band around mean
    ax.fill_between(stats_by_time['timestamp'], 
                     stats_by_time['mean'] - stats_by_time['std'], 
                     stats_by_time['mean'] + stats_by_time['std'],
                     alpha=0.3, color='#2ECC71', label='±1 Std Dev')
    
    # Plot average trend line (main focus)
    ax.plot(stats_by_time['timestamp'], stats_by_time['mean'], 
            color='#E74C3C', linewidth=3, label='Average HR (All Patients)', 
            linestyle='-', marker='o', markersize=4, alpha=0.9)
    
    # Add threshold line
    ax.axhline(y=120, color='#C0392B', linestyle='--', linewidth=3, 
               label='Danger Threshold (120 bpm)', alpha=0.8)
    
    ax.set_xlabel('Timestamp', fontsize=12, fontweight='bold')
    ax.set_ylabel('Heart Rate (bpm)', fontsize=12, fontweight='bold')
    ax.set_title('Heart Rate Trend - Statistical Overview (All Patients)', fontsize=14, fontweight='bold', pad=20)
    ax.legend(loc='upper right', fontsize=10, framealpha=0.9)
    ax.grid(True, alpha=0.3)
    ax.set_ylim(bottom=30, top=140)
    
    # Format x-axis
    ax.xaxis.set_major_formatter(mdates.DateFormatter('%m-%d %H:%M'))
    plt.xticks(rotation=45, ha='right')
    
    plt.tight_layout()
    plt.savefig('visualizations/hr_trend.png', dpi=300, bbox_inches='tight')
    plt.close()
    
    print("  ✓ Saved visualizations/hr_trend.png")
    
    # ----- Chart 2: Oxygen Level Distribution -----
    print("\n[2/3] Generating Oxygen Distribution chart...")
    
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(14, 6))
    
    # Histogram
    ax1.hist(df_vitals['ox'].dropna(), bins=15, color='#3498db', edgecolor='black', alpha=0.7)
    ax1.axvline(x=92, color='red', linestyle='--', linewidth=2, label='Danger Threshold (92)')
    ax1.set_xlabel('Oxygen Saturation (%)', fontsize=12, fontweight='bold')
    ax1.set_ylabel('Frequency', fontsize=12, fontweight='bold')
    ax1.set_title('Oxygen Level Distribution (Histogram)', fontsize=13, fontweight='bold')
    ax1.legend()
    ax1.grid(True, alpha=0.3, axis='y')
    
    # Box plot
    bp = ax2.boxplot(df_vitals['ox'].dropna(), vert=True, patch_artist=True,
                     boxprops=dict(facecolor='#3498db', alpha=0.7),
                     medianprops=dict(color='red', linewidth=2))
    ax2.axhline(y=92, color='red', linestyle='--', linewidth=2, alpha=0.7, label='Danger Threshold (92)')
    ax2.set_ylabel('Oxygen Saturation (%)', fontsize=12, fontweight='bold')
    ax2.set_title('Oxygen Level Distribution (Box Plot)', fontsize=13, fontweight='bold')
    ax2.legend()
    ax2.grid(True, alpha=0.3, axis='y')
    ax2.set_xticklabels(['All Patients'])
    
    plt.suptitle('Oxygen Saturation Analysis', fontsize=14, fontweight='bold', y=1.02)
    plt.tight_layout()
    plt.savefig('visualizations/oxygen_distribution.png', dpi=300, bbox_inches='tight')
    plt.close()
    
    print("  ✓ Saved visualizations/oxygen_distribution.png")
    
    # ----- Chart 3: Anomaly Counts -----
    print("\n[3/3] Generating Anomaly Counts chart...")
    
    if len(df_anomalies) > 0:
        fig, ax = plt.subplots(figsize=(10, 6))
        
        anomaly_counts = df_anomalies['anomaly'].value_counts()
        colors = ['#e74c3c', '#f39c12', '#9b59b6']
        
        bars = ax.bar(anomaly_counts.index, anomaly_counts.values,
                      color=colors[:len(anomaly_counts)], edgecolor='black', linewidth=1.5)
        
        # Add value labels on bars
        for bar in bars:
            height = bar.get_height()
            ax.text(bar.get_x() + bar.get_width()/2., height,
                   f'{int(height)}',
                   ha='center', va='bottom', fontsize=12, fontweight='bold')
        
        ax.set_xlabel('Anomaly Type', fontsize=12, fontweight='bold')
        ax.set_ylabel('Count', fontsize=12, fontweight='bold')
        ax.set_title('Anomaly Detection Summary', fontsize=14, fontweight='bold', pad=20)
        ax.grid(True, alpha=0.3, axis='y')
        
        plt.xticks(rotation=15, ha='right')
        plt.tight_layout()
        plt.savefig('visualizations/anomaly_counts.png', dpi=300, bbox_inches='tight')
        plt.close()
    else:
        # Create empty chart if no anomalies
        fig, ax = plt.subplots(figsize=(10, 6))
        ax.text(0.5, 0.5, 'No Anomalies Detected', 
               ha='center', va='center', fontsize=18, fontweight='bold')
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 1)
        ax.axis('off')
        plt.savefig('visualizations/anomaly_counts.png', dpi=300, bbox_inches='tight')
        plt.close()
    
    print("  ✓ Saved visualizations/anomaly_counts.png")
    
    print("\n✅ All visualizations generated")


# ====================================================================================
# MAIN PIPELINE EXECUTION
# ====================================================================================

def main():
    """Main pipeline orchestration"""
    print("\n")
    print("="*70)
    print(" HOSPITAL HEALTH MONITORING PIPELINE")
    print("="*70)
    print(" Bronze → Silver → Gold Data Processing")
    print("="*70)
    
    # Create directories
    create_directories()
    
    # Step 1: Bronze Layer
    load_to_bronze()
    
    # Step 2: Silver Layer
    df_vitals, df_labs = clean_silver_layer()
    
    # Step 3: Patient Master Table
    patient_master = create_patient_master(df_vitals, df_labs)
    
    # Step 4: Anomaly Detection (Gold Layer)
    df_anomalies = detect_anomalies(df_vitals)
    
    # Step 5: Visualizations
    generate_visualizations(df_vitals, df_anomalies)
    
    # Final summary
    print("\n" + "="*70)
    print(" PIPELINE EXECUTION COMPLETE")
    print("="*70)
    print("\n📊 Summary:")
    print(f"  - Bronze files: 3 (ehr, vitals, labs)")
    print(f"  - Silver files: 3 (clean_vitals, clean_labs, patient_master)")
    print(f"  - Gold files: 1 (anomalies)")
    print(f"  - Visualizations: 3 PNG charts")
    print(f"\n  Total patients: {len(patient_master)}")
    print(f"  Total vitals records: {len(df_vitals)}")
    print(f"  Total anomalies: {len(df_anomalies)}")
    print("\n🎉 Next step: Run 'streamlit run app.py' to view the dashboard!")
    print("="*70 + "\n")


if __name__ == "__main__":
    main()
