from docx import Document

print("VITALS.DOCX CONTENT:")
print("=" * 60)
doc = Document(r'C:\Users\Kashish\Downloads\vitals.docx')
print(f"Number of paragraphs: {len(doc.paragraphs)}")
print(f"Number of tables: {len(doc.tables)}")
print("\nParagraph content (first 30 lines):")
for i, para in enumerate(doc.paragraphs[:30]):
    if para.text.strip():
        print(f"{i}: {para.text}")

if len(doc.tables) > 0:
    print("\nTable found!")
    for table in doc.tables:
        for row in table.rows:
            print([cell.text for cell in row.cells])

print("\n" + "=" * 60)
print("LABS DOCX CONTENT:")
print("=" * 60)
doc = Document(r'C:\Users\Kashish\Downloads\labs- Ajay kumar.docx')
print(f"Number of paragraphs: {len(doc.paragraphs)}")
print(f"Number of tables: {len(doc.tables)}")
print("\nParagraph content (first 30 lines):")
for i, para in enumerate(doc.paragraphs[:30]):
    if para.text.strip():
        print(f"{i}: {para.text}")

if len(doc.tables) > 0:
    print("\nTable found!")
    for table in doc.tables:
        for row in table.rows:
            print([cell.text for cell in row.cells])
